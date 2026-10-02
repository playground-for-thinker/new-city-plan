"""Excel 산출물 — 토지세목조서, 사업비, 토지이용계획, 현금흐름표 등."""
from __future__ import annotations

import io

import geopandas as gpd
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import pipeline
from ..context import analysis
from ..cost import compensation
from ..landuse import plan as landuse_plan
from ..models import Project

_HEAD_FILL = PatternFill("solid", fgColor="DDE5EE")
_THIN = Side(style="thin", color="B0B7C0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _write(writer: pd.ExcelWriter, name: str, blocks: list[tuple[str, pd.DataFrame]]) -> None:
    """한 시트에 (제목, 표) 묶음을 위에서 아래로 쓴다."""
    row = 0
    for title, df in blocks:
        df.to_excel(writer, sheet_name=name, startrow=row + 1, index=False)
        ws = writer.sheets[name]
        ws.cell(row=row + 1, column=1, value=title).font = Font(bold=True, size=12)
        for c in range(1, len(df.columns) + 1):
            cell = ws.cell(row=row + 2, column=c)
            cell.font, cell.fill, cell.border = Font(bold=True), _HEAD_FILL, _BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in range(row + 3, row + 3 + len(df)):
            for c in range(1, len(df.columns) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = _BORDER
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = "#,##0" if float(cell.value).is_integer() else "#,##0.00##"
        row += len(df) + 4
    ws = writer.sheets[name]
    for i, col in enumerate(ws.columns, 1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width * 1.6), 46)


def build_report(project: Project, parcels: gpd.GeoDataFrame, ev: pipeline.Evaluation) -> bytes:
    buf = io.BytesIO()
    m = ev.finance.metrics
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        summary = pd.DataFrame([
            ["사업명", project.name],
            ["사업면적(㎡)", round(ev.total_area)],
            ["편입 필지 수", len(ev.comp.register)],
            ["계획 세대수", round(ev.program.households)],
            ["계획 인구(인)", round(ev.program.population)],
            ["토지보상비(원)", round(ev.comp.land_comp)],
            ["지장물 및 기타 보상비(원)", round(ev.comp.obstacle)],
            ["총사업비 — 조성원가 관점(원)", round(ev.costs.total_cost_basis)],
            ["총사업비 — 예타 재무 분석 관점(원)", round(ev.costs.total_feasibility)],
            ["조성원가(원/㎡, 유상공급 면적 기준)", round(ev.costs.unit_cost)],
            ["분양수입 합계(원)", round(m.total_in)],
            ["재무적 할인율(실질)", project.finance.discount_rate],
            ["PI", round(m.pi, 2) if m.pi == m.pi else None],
            ["FNPV(원)", round(m.fnpv)],
            ["FIRR", round(m.firr, 4) if m.firr is not None else None],
            ["보상배율표 판", project.comp.multiplier_version],
            ["필지 자료 출처", project.parcel_source],
        ], columns=["항목", "값"])
        _write(w, "요약", [("사업 타당성 분석 요약", summary)])

        _write(w, "입지현황", [
            ("지목별 현황", analysis.summary_by(parcels, "jimok")),
            ("용도지역별 현황", analysis.summary_by(parcels, "zone")),
            ("소유구분별 현황", analysis.summary_by(parcels, "owner_group")),
            ("상위계획 검토", pd.DataFrame([{"계획": u.plan, "주요 내용": u.content, "부합 여부": u.conformity, "비고": u.note}
                                       for u in project.upper_plans])),
        ])
        _write(w, "토지세목조서", [("토지세목조서", compensation.register_table(ev.comp))])
        comp_total = pd.DataFrame([
            ["토지보상비(A)", round(ev.comp.land_comp)],
            ["지장물 및 기타 보상비(B)", round(ev.comp.obstacle)],
            ["직접보상비(A+B)", round(ev.comp.direct)],
        ], columns=["구분", "금액(원)"])
        _write(w, "보상비 집계", [
            ("보상비 합계", comp_total),
            ("지목별", compensation.summarize(ev.comp, "jimok")),
            ("용도지역 분류별", compensation.summarize(ev.comp, "zone_class")),
            ("소유구분별", compensation.summarize(ev.comp, "owner")),
        ])
        _write(w, "사업비", [
            ("사업비 구분별 집계", ev.costs.by_category),
            ("사업비 항목별 내역", ev.costs.table.drop(columns=["key", "schedule"])),
        ])
        _write(w, "토지이용계획", [
            (f"토지이용계획표 ({project.active_alt})", landuse_plan.plan_table(project.landuse, ev.total_area).drop(columns="key")),
        ])
        _write(w, "기본구상", [
            ("주택·인구 계획", ev.program.housing.drop(columns="key")),
            ("계획지표", pd.DataFrame([
                ["세대수", round(ev.program.households)], ["인구(인)", round(ev.program.population)],
                ["총밀도(인/ha)", round(ev.program.gross_density, 1)], ["순밀도(인/ha)", round(ev.program.net_density, 1)],
            ], columns=["항목", "값"])),
            ("기반시설 소요", ev.program.infra),
        ])
        _write(w, "분양수입", [("용도별 분양수입", ev.finance.revenue.drop(columns=["key", "major"]))])
        _write(w, "현금흐름", [("연차별 현금흐름 (기준연도 불변가격, 원)", ev.finance.cashflow)])
        _write(w, "민감도", [
            ("변수별 PI 민감도", pipeline.sensitivity(project, ev)),
            ("시나리오 분석", pipeline.scenario_table(project, ev)),
            ("분기점 분석", pipeline.breakeven(project, ev)),
        ])
    return buf.getvalue()
